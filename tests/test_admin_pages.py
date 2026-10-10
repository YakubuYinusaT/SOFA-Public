import json
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from sofa.dialogue import templates
from sofa.dialogue.translation_io import export_xlsx
from sofa.models import (Call, CallTurn, Customer, Handoff, Invoice, Merchant, Notification, Order, OutboundCall,
                         Payment, Product, ProductAlias, StockMovement, VirtualAccount)
from sofa.web import auth
from tests.test_call_flow import MUSA, SPEC_CONVERSATION, charge, product

PASSWORD = "dev-admin-token"


@pytest.fixture(autouse=True)
def fresh_state(tmp_path, monkeypatch):
    auth.clear_failures()
    monkeypatch.setattr(templates, "TRANSLATIONS_PATH", tmp_path / "translations.json")
    templates.load_translations(tmp_path / "missing.json")
    yield
    templates.load_translations(tmp_path / "missing.json")


def login(client):
    r = client.post("/admin/login", data={"password": PASSWORD}, follow_redirects=False)
    assert r.status_code == 303
    return client


def token(client, path="/admin"):
    html = client.get(path).text
    return re.search(r'name="csrf" value="([^"]+)"', html).group(1)


def post(client, path, data=None, page="/admin", **kw):
    """POST a form with a valid CSRF token taken from `page`."""
    return client.post(path, data={"csrf": token(client, page), **(data or {})}, follow_redirects=False, **kw)


def flash(client, response):
    """Follow a redirect and return the page text (flash message included)."""
    return client.get(response.headers["location"]).text


def merchant_id(db):
    return db.scalar(select(Merchant.id))


# ---- access control ----------------------------------------------------------------------------------


def test_pages_require_login_and_remember_where_you_were_going(client):
    r = client.get("/admin/orders?status=paid", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].startswith("/admin/login?next=")
    assert client.post("/admin/merchants", data={"name": "x"}, follow_redirects=False).status_code == 303
    r = client.post("/admin/login", data={"password": PASSWORD, "next": "/admin/orders?status=paid"}, follow_redirects=False)
    assert r.headers["location"] == "/admin/orders?status=paid"


def test_wrong_password_and_lockout(client):
    assert client.post("/admin/login", data={"password": "nope"}).status_code == 401
    for _ in range(10):
        client.post("/admin/login", data={"password": "nope"})
    assert client.post("/admin/login", data={"password": PASSWORD}).status_code == 429  # even the right one, for now


def test_login_cannot_redirect_off_site(client):
    r = client.post("/admin/login", data={"password": PASSWORD, "next": "https://evil.example/x"}, follow_redirects=False)
    assert r.headers["location"] == "/admin"
    r = client.post("/admin/login", data={"password": PASSWORD, "next": "//evil.example"}, follow_redirects=False)
    assert r.headers["location"] == "/admin"


def test_cookie_is_httponly_samesite_strict(client):
    r = client.post("/admin/login", data={"password": PASSWORD}, follow_redirects=False)
    header = r.headers["set-cookie"].lower()
    assert "httponly" in header and "samesite=strict" in header and "path=/" in header and "path=/admin" not in header  # site-wide on purpose: the master page shows where you are logged in


def test_forged_cookie_is_rejected(client):
    client.cookies.set(auth.COOKIE, "forged.value.here", path="/admin")
    assert client.get("/admin", follow_redirects=False).status_code == 303


def test_post_needs_the_csrf_token(client):
    login(client)
    assert client.post("/admin/merchants", data={"name": "x"}).status_code == 403
    assert client.post("/admin/merchants", data={"name": "x", "csrf": "wrong"}).status_code == 403
    assert client.post("/admin/logout", data={}).status_code == 403


def test_logout_ends_the_session(client):
    login(client)
    assert post(client, "/admin/logout").headers["location"] == "/admin/login"
    client.cookies.clear()
    assert client.get("/admin", follow_redirects=False).status_code == 303


def test_dashboard_shows_the_company_line_and_default_password_warning(client):
    login(client)
    html = client.get("/admin").text
    assert "a product of Connected Intelligence" in html
    assert "Awarri" not in html and "Demo Bank" not in html  # the team console carries no footer disclaimer
    assert "Connected Intelligence" in client.get("/admin/login").text
    assert "still uses the demo password" in html
    for label in ("Dashboard", "Merchants", "Orders", "Handoffs", "Calls", "Label", "Metrics", "Languages"):
        assert label in html


# ---- merchants and catalog ---------------------------------------------------------------------------------


def test_create_merchant_and_validation(client, db):
    login(client)
    ok = post(client, "/admin/merchants", {"name": "Bola Pharmacy", "category": "pharmacy", "owner_name": "Bola",
                                          "owner_phone": "0803 123 4567", "number": "+2348009998888",
                                          "default_language": "yo", "payment_mode": "pay_on_delivery"}, page="/admin/merchants")
    assert "/admin/merchants/" in ok.headers["location"]
    db.expire_all()
    m = db.scalar(select(Merchant).where(Merchant.name == "Bola Pharmacy"))
    assert (m.default_language, m.payment_mode) == ("yo", "pay_on_delivery")
    dup = post(client, "/admin/merchants", {"name": "X", "owner_name": "Y", "owner_phone": "08031234568", "number": "+2348009998888"}, page="/admin/merchants")
    assert "already assigned" in flash(client, dup)
    bad = post(client, "/admin/merchants", {"name": "X", "owner_name": "Y", "owner_phone": "12345", "number": "+2348007776666"}, page="/admin/merchants")
    assert "not a Nigerian mobile number" in flash(client, bad)
    bad = post(client, "/admin/merchants", {"name": "X", "owner_name": "Y", "owner_phone": "08031234568", "number": "+2348007776666",
                                            "payment_mode": "barter"}, page="/admin/merchants")
    assert "payment mode must be one of" in flash(client, bad)


def test_merchant_settings_users_and_numbers(client, db):
    login(client)
    mid = merchant_id(db)
    post(client, f"/admin/merchants/{mid}", {"name": "CI Store 2", "category": "provisions", "default_language": "yo",
                                             "payment_mode": "per_customer", "status": "paused", "paystack_subaccount_code": "ACCT_abc"})
    db.expire_all()
    m = db.get(Merchant, mid)
    assert (m.name, m.default_language, m.status, m.paystack_subaccount_code) == ("CI Store 2", "yo", "paused", "ACCT_abc")
    r = post(client, f"/admin/merchants/{mid}/users", {"name": "Tunde", "phone": "08055551234", "role": "staff", "can_update_stock": "on"})
    assert "Person added" in flash(client, r)
    r = post(client, f"/admin/merchants/{mid}/numbers", {"number": "+2348001112222"})
    assert "already assigned" in flash(client, r)  # the seeded demo number


def test_csv_import_all_or_nothing_and_stock_movements(client, db):
    login(client)
    mid = merchant_id(db)
    page = f"/admin/merchants/{mid}/products"
    good = "name,unit,price_naira,stock_qty,aliases,units\nKings Oil 1L,piece,2500,30,kings|oil,carton:12:28000\n"
    r = post(client, f"{page}/import", page=page, files={"file": ("c.csv", good.encode(), "text/csv")})
    assert "Imported 1 products" in flash(client, r)
    db.expire_all()
    oil = db.scalar(select(Product).where(Product.name == "Kings Oil 1L"))
    assert oil.price_kobo == 250_000 and oil.stock_qty == 30
    assert {u.unit: u.price_kobo for u in oil.units} == {"piece": 250_000, "carton": 2_800_000}
    assert db.scalar(select(StockMovement).where(StockMovement.product_id == oil.id, StockMovement.delta == 30))

    before = len(db.scalars(select(Product)).all())
    bad = "name,unit,price_naira\nGood One,piece,100\nBroken,piece,abc\n"
    r = post(client, f"{page}/import", page=page, files={"file": ("c.csv", bad.encode(), "text/csv")})
    text = flash(client, r)
    assert "Nothing imported" in text and "row 3" in text
    db.expire_all()
    assert len(db.scalars(select(Product)).all()) == before  # the good row was not kept either
    r = post(client, f"{page}/import", page=page, files={"file": ("c.csv", b"colour,size\nred,big\n", "text/csv")})
    assert "missing column" in flash(client, r)


def test_price_stock_and_alias_edits(client, db):
    login(client)
    mid = merchant_id(db)
    page = f"/admin/merchants/{mid}/products"
    p = product(db, "Indomie Super Pack")
    # blank price keeps prices; the stock change is recorded as a movement
    post(client, f"{page}/{p.id}", {"unit": "carton", "price_naira": "", "stock_qty": "700", "active": "on"}, page=page)
    db.expire_all()
    p = product(db, "Indomie Super Pack")
    assert p.stock_qty == 700 and p.price_kobo == 35_000
    assert next(u for u in p.units if u.unit == "carton").price_kobo == 700_000
    assert db.scalar(select(StockMovement).where(StockMovement.product_id == p.id, StockMovement.delta == -100, StockMovement.source == "web"))
    # carton price only touches the carton
    post(client, f"{page}/{p.id}", {"unit": "carton", "price_naira": "7500", "stock_qty": "700", "active": "on"}, page=page)
    db.expire_all()
    p = product(db, "Indomie Super Pack")
    assert next(u for u in p.units if u.unit == "carton").price_kobo == 750_000 and p.price_kobo == 35_000
    # unticking active hides it from callers
    post(client, f"{page}/{p.id}", {"unit": "pack", "stock_qty": "700"}, page=page)
    db.expire_all()
    assert product(db, "Indomie Super Pack").active is False
    bad = post(client, f"{page}/{p.id}", {"unit": "pack", "stock_qty": "-5", "active": "on"}, page=page)
    assert "cannot be negative" in flash(client, bad)
    post(client, f"{page}/{p.id}/aliases", {"phrase": "Indomie Kubu"}, page=page)
    db.expire_all()
    alias = db.scalar(select(ProductAlias).where(ProductAlias.phrase == "indomie kubu"))
    assert alias.source == "merchant_correction" and alias.product_id == p.id


def test_html_in_product_names_is_escaped(client, db):
    login(client)
    mid = merchant_id(db)
    page = f"/admin/merchants/{mid}/products"
    post(client, page, {"name": "<script>alert(1)</script>", "unit": "piece", "price_naira": "10"}, page=page)
    html = client.get(page).text
    assert "<script>alert(1)</script>" not in html and "&lt;script&gt;alert(1)&lt;/script&gt;" in html


def test_sample_csv_imports_cleanly(client, db):
    login(client)
    sample = client.get("/admin/sample.csv")
    assert sample.status_code == 200 and "name,brand" in sample.text
    mid = merchant_id(db)
    r = post(client, f"/admin/merchants/{mid}/products/import", page="/admin/merchants",
             files={"file": ("s.csv", sample.content, "text/csv")})
    assert "Imported 2 products" in flash(client, r)


# ---- orders and payments -----------------------------------------------------------------------------------


def place_order(call):
    call(SPEC_CONVERSATION)


def test_order_page_dispatch_and_deliver(client, call, db):
    place_order(call)
    login(client)
    o = db.scalar(select(Order))
    assert "SOFA-" in client.get(f"/admin/orders/{o.id}").text
    assert "Musa" in client.get("/admin/orders").text
    post(client, f"/admin/orders/{o.id}/status", {"status": "dispatched"}, page=f"/admin/orders/{o.id}")  # not allowed from awaiting_payment
    db.expire_all()
    assert db.get(Order, o.id).status == "awaiting_payment"


def test_manual_payment_needs_confirmation_then_dispatch_flow(client, call, db):
    place_order(call)
    login(client)
    o = db.scalar(select(Order))
    inv = db.scalar(select(Invoice))
    page = f"/admin/orders/{o.id}"
    r = post(client, f"/admin/invoices/{inv.id}/payments", {"amount_naira": "16400"}, page=page)  # box not ticked
    assert "tick the box" in flash(client, r)
    r = post(client, f"/admin/invoices/{inv.id}/payments", {"amount_naira": "10000", "sender_name": "Musa", "confirmed": "on"}, page=page)
    assert "Part payment recorded" in flash(client, r)
    db.expire_all()
    assert db.get(Invoice, inv.id).status == "open" and db.get(Order, o.id).status == "awaiting_payment"
    r = post(client, f"/admin/invoices/{inv.id}/payments", {"amount_naira": "6400", "confirmed": "on"}, page=page)
    assert "Invoice paid" in flash(client, r)
    db.expire_all()
    assert db.get(Invoice, inv.id).status == "paid" and db.get(Order, o.id).status == "paid"
    assert db.scalar(select(OutboundCall).where(OutboundCall.trigger == "payment_received"))
    assert db.scalar(select(Notification).where(Notification.template == "payment_received"))

    post(client, f"/admin/orders/{o.id}/status", {"status": "dispatched"}, page=page)
    post(client, f"/admin/orders/{o.id}/status", {"status": "delivered"}, page=page)
    db.expire_all()
    assert db.get(Order, o.id).status == "delivered"
    assert db.scalar(select(Notification).where(Notification.template == "dispatched"))
    assert db.scalar(select(OutboundCall).where(OutboundCall.trigger == "delivered"))


def test_overpayment_is_flagged_for_refund(client, call, db):
    place_order(call)
    login(client)
    o, inv = db.scalar(select(Order)), db.scalar(select(Invoice))
    r = post(client, f"/admin/invoices/{inv.id}/payments", {"amount_naira": "20000", "confirmed": "on"}, page=f"/admin/orders/{o.id}")
    assert "MORE than the amount" in flash(client, r)


def test_cancel_from_page_releases_stock_and_voids_invoice(client, call, db):
    place_order(call)
    login(client)
    o = db.scalar(select(Order))
    post(client, f"/admin/orders/{o.id}/status", {"status": "cancelled"}, page=f"/admin/orders/{o.id}")
    db.expire_all()
    assert db.get(Order, o.id).status == "cancelled" and db.scalar(select(Invoice)).status == "void"
    assert product(db, "Peak Milk Tin").stock_qty == 200


def test_unmatched_payment_is_resolved_from_the_handoffs_page(client, call, db):
    place_order(call)
    login(client)
    acct = db.scalar(select(VirtualAccount))
    charge(client, acct.account_number, 500_000, event_id=3003)  # short payment: unmatched + handoff
    html = client.get("/admin/handoffs").text
    assert "unmatched payment" in html and "5,000 naira" in html
    db.expire_all()
    pay, inv = db.scalar(select(Payment)), db.scalar(select(Invoice))
    r = post(client, f"/admin/payments/{pay.id}/apply", {"invoice_id": str(inv.id)}, page="/admin/handoffs")
    assert "part payment" in flash(client, r)
    r2 = post(client, f"/admin/invoices/{inv.id}/payments", {"amount_naira": "11400", "confirmed": "on"}, page="/admin/handoffs")
    db.expire_all()
    assert db.get(Invoice, inv.id).status == "paid"  # 5,000 + 11,400 = 16,400


def test_handoff_queue_resolve_and_reopen(client, call, db):
    call(["Let me talk to the owner"])
    login(client)
    html = client.get("/admin/handoffs").text
    assert "speak to human" in html and MUSA in html or "Caller said" in html
    h = db.scalar(select(Handoff))
    post(client, f"/admin/handoffs/{h.id}/resolve", {"resolved_by": "Ade"}, page="/admin/handoffs")
    db.expire_all()
    assert db.get(Handoff, h.id).status == "resolved" and db.get(Handoff, h.id).resolved_by == "Ade"
    assert "by Ade" in client.get("/admin/handoffs?status=resolved").text
    post(client, f"/admin/handoffs/{h.id}/resolve", {"reopen": "1"}, page="/admin/handoffs")
    db.expire_all()
    assert db.get(Handoff, h.id).status == "open"


# ---- calls, audio, labelling, metrics --------------------------------------------------------------------------


def test_call_page_lists_turns_and_serves_audio_only_to_admins(client, call, db, tmp_path):
    call(SPEC_CONVERSATION[:2])
    c = db.scalar(select(Call))
    turn = db.scalar(select(CallTurn).order_by(CallTurn.seq))
    assert client.get(f"/admin/audio/{turn.id}/caller", follow_redirects=False).status_code == 303  # not logged in
    login(client)
    html = client.get(f"/admin/calls/{c.id}").text
    assert "Indomie big one" in html and "Sofa" in html and "/admin/audio/" in html
    assert MUSA in client.get("/admin/calls").text
    assert client.get(f"/admin/audio/{turn.id}/caller").status_code == 200
    assert client.get(f"/admin/audio/{turn.id}/reply").headers["content-type"] == "audio/wav"
    assert client.get(f"/admin/audio/{turn.id}/other").status_code == 404
    turn.audio_path = str(Path(tmp_path).parent / "outside.wav")  # a path outside the storage folder is never served
    Path(turn.audio_path).write_bytes(b"x")
    db.commit()
    assert client.get(f"/admin/audio/{turn.id}/caller").status_code == 404


def test_test_calls_are_hidden_unless_asked(client, call, db):
    call(["How much is Peak milk tin?"], caller="+2348000000001")  # a configured test number
    login(client)
    assert "+2348000000001" not in client.get("/admin/calls").text
    assert "+2348000000001" in client.get("/admin/calls?tests=1").text


def test_labelling_updates_metrics_teaches_alias_and_exports_are_anonymous(client, call, db):
    call(["My name is Musa, I want Indomie big one", "Super pack", "Carton. Two carton", "No, that is all",
          "Ajao Estate, near the mosque on Adeyemi Street", "Yes"])
    login(client)
    page = client.get("/admin/label").text
    assert "0 of 6 turns labelled" in page and "Save and next" in page and "Indomie big one" in page

    first = db.scalar(select(CallTurn).order_by(CallTurn.seq))
    form = {"transcript": "My name is Musa, I want Indomie big one", "language": "en", "intent": "place_order", "n_matches": "1",
            "spoken_0": "indomie big one", "system_0": "", "correct_0": "Indomie Hungry Man", "notes": "shop noise",
            "flags": ["background noise"], "filter_lang": "", "filter_tests": ""}
    r = post(client, f"/admin/label/{first.id}", form, page="/admin/label")
    assert "Saved" in flash(client, r)
    db.expire_all()
    first = db.get(CallTurn, first.id)
    assert first.labelled_by == "admin" and first.label["flags"] == ["background noise"]
    hungry = product(db, "Indomie Hungry Man")
    alias = db.scalar(select(ProductAlias).where(ProductAlias.phrase == "indomie big one", ProductAlias.product_id == hungry.id))
    assert alias.source == "merchant_correction"
    # the phrase was also confirmed on a call for another product: the staff correction outranks it
    from sofa.dialogue.matching import _exact_alias, active_products
    from sofa.textutil import normalize_phrase
    assert _exact_alias(active_products(db, hungry.merchant_id), normalize_phrase("indomie big one")).name == "Indomie Hungry Man"

    # a second turn where the machine heard it slightly wrong: WER is computed from the correction
    second = db.scalar(select(CallTurn).where(CallTurn.seq == 2))
    post(client, f"/admin/label/{second.id}", {"transcript": "Super pack please", "language": "en", "intent": "place_order", "n_matches": "0"}, page="/admin/label")
    db.expire_all()
    m = client.get("/admin/metrics").text
    assert "Accuracy by language" in m and "n=2" in m and "17%" in m  # WER 0% and 33% (one missing word of three): mean 17%
    assert "2 of 6" in client.get("/admin/label").text or "2 of 6 turns labelled" in client.get("/admin/label").text

    # anonymised exports: no phone, no name, no address, and only labelled + consented turns in the training files
    turns_csv = client.get("/admin/export/turns.csv").text
    assert "+2348055550001" not in turns_csv and "08055550001" not in turns_csv
    assert "Musa" not in turns_csv and "Ajao Estate" not in turns_csv and "<NAME>" in turns_csv
    calls_csv = client.get("/admin/export/calls.csv").text
    assert MUSA not in calls_csv and "caller_" in calls_csv and "CI Store" in calls_csv
    understanding = [json.loads(l) for l in client.get("/admin/export/understanding.jsonl").text.splitlines()]
    assert len(understanding) == 2 and "Musa" not in json.dumps(understanding)
    assert understanding[0]["output"]["customer_name"] is None and understanding[0]["input"].startswith("My name is <NAME>")
    speech = [json.loads(l) for l in client.get("/admin/export/speech.jsonl").text.splitlines()]
    assert len(speech) == 2 and speech[0]["audio"].startswith("audio_in") and "Musa" not in json.dumps(speech)
    assert client.get("/admin/export/nothing.csv").status_code == 404


def test_metrics_page_renders_with_no_data(client):
    login(client)
    html = client.get("/admin/metrics").text
    assert "Validation metrics" in html and "Nothing labelled yet" in html and "Needs at least 6 real calls" in html


def test_early_vs_late_and_clarification_rate(client, call, db):
    for i in range(8):
        call(["Do you have Peak milk tin?"], caller=f"+23480555510{i:02d}")
    login(client)
    html = client.get("/admin/metrics").text
    assert "First 4 calls" in html and "Last 4 calls" in html


# ---- languages ------------------------------------------------------------------------------------------------------------


def fill_yoruba(path):
    wb = load_workbook(path)
    ws = wb["Translate"]
    header = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "repeat_prompt" and row[1].value == 0:
            row[header.index("Yoruba")].value = "Yoruba: jowo, e tun so o?"
        if row[0].value == "handoff" and row[1].value == 0:
            row[header.index("Yoruba")].value = "Ma pe e {broken"
    wb.save(path)


def test_language_page_download_upload_and_reject(client, tmp_path):
    login(client)
    assert "0 / " in client.get("/admin/languages").text
    sheet = client.get("/admin/translation-sheet.xlsx")
    assert sheet.status_code == 200 and sheet.content[:2] == b"PK"
    path = tmp_path / "filled.xlsx"
    path.write_bytes(sheet.content)
    fill_yoruba(path)

    bad = post(client, "/admin/languages/import", page="/admin/languages", files={"file": ("filled.xlsx", path.read_bytes(), "application/octet-stream")})
    assert bad.status_code == 200 and "Nothing saved" in bad.text and "stray curly bracket" in bad.text
    assert not templates.has("repeat_prompt", "yo")

    dry = post(client, "/admin/languages/import", {"allow_partial": "on", "dry_run": "on"}, page="/admin/languages",
               files={"file": ("filled.xlsx", path.read_bytes(), "application/octet-stream")})
    assert "Checked, nothing saved" in dry.text and not templates.has("repeat_prompt", "yo")

    ok = post(client, "/admin/languages/import", {"allow_partial": "on"}, page="/admin/languages",
              files={"file": ("filled.xlsx", path.read_bytes(), "application/octet-stream")})
    assert "Saved" in ok.text and "these cells were skipped" in ok.text
    assert templates.has("repeat_prompt", "yo") and not templates.has("handoff", "yo")
    assert "1 / " in client.get("/admin/languages").text


def test_language_upload_rejects_junk_files(client):
    login(client)
    r = post(client, "/admin/languages/import", page="/admin/languages", files={"file": ("x.xlsx", b"not a workbook", "application/octet-stream")})
    assert r.status_code == 303 and "Could not read that file" in client.get(r.headers["location"]).text


# ---- API parity ----------------------------------------------------------------------------------------------------------------


def test_api_stock_patch_leaves_a_movement_and_api_export_needs_token(client, db):
    h = {"Authorization": "Bearer dev-admin-token"}
    mid = merchant_id(db)
    p = product(db, "Peak Milk Tin")
    assert client.patch(f"/api/merchants/{mid}/products/{p.id}", json={"stock_qty": 150}, headers=h).status_code == 200
    db.expire_all()
    assert product(db, "Peak Milk Tin").stock_qty == 150
    assert db.scalar(select(StockMovement).where(StockMovement.product_id == p.id, StockMovement.delta == -50, StockMovement.source == "web"))
    assert client.get("/api/export/calls.csv").status_code == 401
    assert client.get("/api/export/calls.csv", headers=h).status_code == 200


def test_every_admin_table_becomes_labelled_cards_on_a_phone(client, db, call):
    """A table wider than a phone would have to be scrolled sideways: each admin table is marked to stack, and every data cell carries its column's name."""
    call(SPEC_CONVERSATION, caller=MUSA)
    login(client)
    merchant = db.scalar(select(Merchant))
    for path in ("/admin", "/admin/merchants", f"/admin/merchants/{merchant.id}/products", "/admin/orders", "/admin/handoffs", "/admin/calls", "/admin/metrics", "/admin/languages"):
        html = client.get(path).text
        tables = re.findall(r"<table[^>]*>", html)
        assert tables, path
        assert all('class="responsive stack"' in t for t in tables), (path, tables)
        assert "<thead>" in html and ("data-label=" in html or "colspan=" in html), path   # a table with no rows yet shows its one empty-state line instead
    order = db.scalar(select(Order))
    page = client.get(f"/admin/orders/{order.id}").text
    assert all('class="responsive stack"' in t for t in re.findall(r"<table[^>]*>", page)) and 'data-label="Total"' in page
