import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from sofa.audio import AudioService
from sofa.dialogue import templates
from sofa.dialogue.translation_io import HEADERS, build, check_cell, export_xlsx, import_sheet, read_sheet
from sofa.models import CallTurn


@pytest.fixture(autouse=True)
def isolated_translations(tmp_path, monkeypatch):
    """Never touch the real translations.json; restore an empty (English-only) state afterwards."""
    monkeypatch.setattr(templates, "TRANSLATIONS_PATH", tmp_path / "translations.json")
    templates.load_translations(tmp_path / "missing.json")
    yield
    templates.load_translations(tmp_path / "missing.json")


def fill(path: Path, cells: dict[tuple[str, int, str], str]):
    """cells: (key, variant, column header) -> text, written into the exported workbook."""
    wb = load_workbook(path)
    ws = wb["Translate"]
    header = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        k, v = row[0].value, row[1].value
        for (ck, cv, col), text in cells.items():
            if (k, v) == (ck, cv):
                row[header.index(col)].value = text
    wb.save(path)


def test_export_lists_every_translatable_phrase_most_used_first(tmp_path):
    out = tmp_path / "sheet.xlsx"
    rows = export_xlsx(out)
    ws = load_workbook(out)["Translate"]
    assert [c.value for c in ws[1]] == HEADERS
    assert ws.max_row - 1 == rows
    keys = [r[0].value for r in ws.iter_rows(min_row=2)]
    assert set(keys) == set(templates.translatable_keys())
    assert "language_menu" not in keys and "no_service" not in keys  # always English
    first_p2 = next(i for i, k in enumerate(keys) if k not in templates.P1)
    assert all(k in templates.P1 for k in keys[:first_p2]) and keys[0] in templates.P1
    readback = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == "readback"]
    assert len(readback) == 3 and readback[0][2].startswith("Required") and readback[2][2].startswith("Optional extra")
    assert "{items} =" in readback[0][5] and "{address} =" in readback[0][5]  # translators see what fills each slot
    assert "READ-BACK" in readback[0][3]


def test_every_english_phrase_has_context_and_placeholder_help():
    for key, phrasings in templates.EN.items():
        assert key in templates.CONTEXT, f"no translator context for {key}"
        for p in phrasings:
            for name in {ph.strip("{}").split(".")[0] for ph in templates.placeholders(p)}:
                assert name in templates.PLACEHOLDER_HELP, f"{key} uses {{{name}}} with no explanation"


def test_round_trip_import_and_runtime_use(tmp_path):
    out = tmp_path / "sheet.xlsx"
    export_xlsx(out)
    fill(out, {
        ("repeat_prompt", 0, "Yoruba"): "Yoruba: jowo, e tun so o?",
        ("ask_address", 0, "Yoruba"): "Yoruba: nibo ni ki a gbe e de?",
        ("ack", 0, "Yoruba"): "Sho", ("ack", 1, "Yoruba"): "Ti gbo",
        ("added", 0, "Hausa"): "Hausa: {ack}. {qty_unit} na {product}, {line_total}. Akwai wani abu?",
    })
    res = import_sheet(out)
    assert not res.errors
    saved = json.loads((tmp_path / "translations.json").read_text(encoding="utf-8"))
    assert saved["yo"]["repeat_prompt"] == ["Yoruba: jowo, e tun so o?"] and saved["yo"]["ack"] == ["Sho", "Ti gbo"]
    assert "ig" not in saved
    assert templates.pick("repeat_prompt", "yo") == "Yoruba: jowo, e tun so o?"
    assert templates.acks("yo") == ["Sho", "Ti gbo"] and templates.acks("ha") == templates.EN["ack"]
    done, total, missing = templates.coverage("yo")
    assert done == 3 and "added" in missing and total == len(templates.translatable_keys())


def test_bad_cells_are_reported_and_block_the_import(tmp_path):
    out = tmp_path / "sheet.xlsx"
    export_xlsx(out)
    fill(out, {
        ("readback", 0, "Yoruba"): "Eyi ni: {items}. Se ki n gbe e?",                       # missing {total} {address}
        ("added", 0, "Hausa"): "Na kara guda 2 na {product}, {line_total}. Akwai wani abu?",  # digit, missing placeholders
        ("ask_address", 0, "Igbo"): "Ebee ka anyi ga-ebu ya",                               # not a question
        ("handoff", 0, "Yoruba"): "Ma pe e {pada",                                          # stray brace
        ("greet_new", 0, "Yoruba"): "{greeting}, Sofa ni mi lati {merchant}. Tani e, kini e fe?",  # good
    })
    res = import_sheet(out)
    joined = "\n".join(res.errors)
    assert "Yoruba / readback / variant 0" in joined and "missing {address} {total}" in joined
    assert "contains a number" in joined and "Hausa / added" in joined
    assert "must end with ?" in joined and "stray curly bracket" in joined
    assert not (tmp_path / "translations.json").exists()  # nothing written on errors

    res = import_sheet(out, allow_partial=True)  # keeps only the valid cell
    assert res.translations["yo"] == {"greet_new": ["{greeting}, Sofa ni mi lati {merchant}. Tani e, kini e fe?"]}
    assert templates.has("greet_new", "yo") and not templates.has("readback", "yo")


def test_check_cell_rules():
    assert check_cell("filler", 0, "Duro die.") == []
    assert any("placeholders must match" in p for p in check_cell("greet_owner", 0, "Bawo {naam}."))
    assert not any("number" in p for p in check_cell("greet_owner", 0, "Bawo {name}. Fun apere ni 20 paali."))  # English has digits too
    assert any("identical to the English" in p for p in check_cell("added", 0, templates.EN["added"][0]))
    assert any("placeholders not allowed" in p for p in check_cell("ask_what", 2, "Bawo {merchant}?"))  # extra wording


def test_csv_sheet_also_imports(tmp_path):
    csv_path = tmp_path / "sheet.csv"
    csv_path.write_text(
        "key,variant,status,when it is spoken,English,what {placeholders} will contain,Yoruba,Hausa,Igbo,notes\n"
        'still_there,0,Required,,Are you still there?,,Yoruba: o wa nibe?,,,\n', encoding="utf-8-sig")
    res = build(read_sheet(csv_path))
    assert not res.errors and res.translations["yo"]["still_there"] == ["Yoruba: o wa nibe?"]


def test_untranslated_phrases_are_spoken_in_english_voice(call, db, tmp_path):
    """A Yoruba caller with only some phrases translated: translated ones use the Yoruba voice, the
    rest fall back to English text AND the English voice (never English words in a Yoruba voice)."""
    out = tmp_path / "sheet.xlsx"
    export_xlsx(out)
    fill(out, {("repeat_prompt", 0, "Yoruba"): "Yoruba: jowo, e tun so o?"})
    assert not import_sheet(out).errors
    call(["blah blah zzz", "Do you have Peak milk tin?"], lang="yo")
    turns = list(db.scalars(select(CallTurn).order_by(CallTurn.seq)))
    assert turns[0].reply_text == "Yoruba: jowo, e tun so o?"
    assert Path(turns[0].reply_audio_path).stem == AudioService.key(turns[0].reply_text, "yo")
    assert "Yes, we have Peak Milk Tin" in turns[1].reply_text  # untranslated: English text
    assert Path(turns[1].reply_audio_path).stem == AudioService.key(turns[1].reply_text, "en")


def test_translated_acks_and_lists_are_used(call, db, tmp_path):
    out = tmp_path / "sheet.xlsx"
    export_xlsx(out)
    fill(out, {
        ("ack", 0, "Yoruba"): "Sho", ("ack", 1, "Yoruba"): "Ti gbo",
        ("word_or", 0, "Yoruba"): "tabi",
        ("clarify_variant", 0, "Yoruba"): "Yoruba: a ni {options}. Eyi wo?",
    })
    assert not import_sheet(out).errors
    call(["I want Peak milk"], lang="yo")
    reply = db.scalar(select(CallTurn).order_by(CallTurn.seq.desc())).reply_text
    assert reply.startswith("Yoruba: a ni Peak Milk Tin at 800 naira tabi Peak Milk Sachet at 100 naira")
