"""The pilot runs English and Yoruba. Hausa and Igbo stay in the code, switched off by ENABLED_LANGUAGES."""

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from scripts.seed import DEMO_NUMBER, seed
from scripts.simulate_call import run_call
from sofa.audio import AudioService, mock_recording_url
from sofa.config import Settings
from sofa.dialogue import language, templates
from sofa.dialogue.translation_io import export_xlsx
from sofa.main import create_app
from sofa.models import Customer
from tests.conftest import SECRET
from tests.test_admin_pages import flash, login, post

CALLER = "+2348055550001"


@pytest.fixture(autouse=True)
def english_only_templates(tmp_path, monkeypatch):
    monkeypatch.setattr(templates, "TRANSLATIONS_PATH", tmp_path / "translations.json")
    templates.load_translations(tmp_path / "missing.json")
    yield
    templates.load_translations(tmp_path / "missing.json")


def app_with(settings: Settings, enabled: str):
    settings.enabled_languages = enabled
    app = create_app(settings)
    seed(app.state.svc.session_factory)
    asked: list[str] = []
    real = app.state.svc.asr.transcribe

    async def spy(audio, language):
        asked.append(language)
        return await real(audio, language)

    app.state.svc.asr.transcribe = spy
    return app, asked


def phone(client, lines, lang="en", caller=CALLER):
    return run_call(client, lines, caller, DEMO_NUMBER, lang, SECRET, echo=lambda *_: None)


# ---- the setting ---------------------------------------------------------------------------------------


def test_default_is_english_and_yoruba_and_english_is_always_on():
    assert Settings(_env_file=None).languages == ("en", "yo")
    assert Settings(enabled_languages="yo", _env_file=None).languages == ("en", "yo")  # English cannot be switched off
    assert Settings(enabled_languages="ig, yo ,xx,ha", _env_file=None).languages == ("en", "ig", "yo", "ha")  # order = menu order
    assert Settings(enabled_languages="", _env_file=None).languages == ("en",)


# ---- what a caller experiences ---------------------------------------------------------------------------------


def test_only_english_and_yoruba_models_are_asked(settings):
    app, asked = app_with(settings, "en,yo")
    with TestClient(app) as client:
        phone(client, ["How much is Peak milk tin?"], lang="yo")
    assert sorted(asked) == ["en", "yo"]  # two models, not four


def test_all_four_are_asked_when_all_are_enabled(settings):
    app, asked = app_with(settings, "en,yo,ha,ig")
    with TestClient(app) as client:
        phone(client, ["How much is Peak milk tin?"], lang="ha")
    assert sorted(asked) == ["en", "ha", "ig", "yo"]


def test_english_only_skips_detection_completely(settings):
    app, asked = app_with(settings, "en")
    with TestClient(app) as client:
        phone(client, ["How much is Peak milk tin?"], lang="en")
    assert asked == ["en"]


def raw_turn(client, session, who, text_tagged):
    """One caller turn straight against the voice endpoint, returning the Voice XML."""
    return client.post(f"/voice/turn/{SECRET}", data={"sessionId": session, "callerNumber": who, "destinationNumber": DEMO_NUMBER,
                                                        "recordingUrl": mock_recording_url(text_tagged[1], text_tagged[0])}).text


def start(client, session, who):
    client.post(f"/voice/inbound/{SECRET}", data={"sessionId": session, "callerNumber": who, "destinationNumber": DEMO_NUMBER})


def test_spoken_answers_in_many_spellings_are_understood():
    both = ("en", "yo")
    for said, want in [("Yoruba", "yo"), ("youruba", "yo"), ("Yorùbá", "yo"), ("Yoruba please", "yo"), ("ede Yoruba", "yo"),
                       ("English", "en"), ("inglish", "en"), ("English please", "en"), ("Gẹ̀ẹ́sì", "en"), ("yes english is fine", "en")]:
        assert language.spoken_choice([said], both) == want, said
    for unclear in ["", "hello", "English or Yoruba", "blah blah"]:  # nothing named, or two named: not an answer
        assert language.spoken_choice([unclear], both) is None, unclear
    assert language.spoken_choice(["hausa"], both) is None  # not offered on this line
    assert language.spoken_choice(["hausa"], ("en", "yo", "ha")) == "ha"
    assert language.spoken_choice(["nonsense", "Yoruba"], both) == "yo"  # any model's transcript may carry the word


def test_an_unrecognised_language_is_asked_about_out_loud_never_with_a_menu(settings):
    app, asked = app_with(settings, "en,yo")
    with TestClient(app) as client:
        start(client, "s1", CALLER)
        xml = raw_turn(client, "s1", CALLER, ("ha", "Ina son madara"))
        assert "<Record" in xml and "GetDigits" not in xml  # an ordinary turn: the caller answers by speaking
        replies = phone(client, ["[ha] Ina son madara"])
        assert replies[-1] == "Sorry, I didn't catch your language. Would you like to speak English or Yoruba?"
        assert "press" not in replies[-1].lower()


def test_caller_answers_by_voice_and_the_conversation_continues_in_that_language(settings):
    app, asked = app_with(settings, "en,yo")
    with TestClient(app) as client:
        replies = phone(client, ["[ha] Ina son madara", "[yo] Yoruba please", "[yo] How much is Peak milk tin?"])
        assert replies[1].startswith("Sorry, I didn't catch your language")
        assert replies[2] == "Thank you. What would you like today?"  # no Yoruba wording yet: English text, English voice
        assert "800 naira per tin" in replies[3]
        with app.state.svc.session_factory() as db:
            assert db.scalar(select(Customer).where(Customer.phone == CALLER)).language == "yo"  # remembered for next time


def test_saying_english_also_works(settings):
    app, _ = app_with(settings, "en,yo")
    with TestClient(app) as client:
        replies = phone(client, ["[ha] Ina son madara", "[en] English please", "[en] How much is Peak milk tin?"])
        assert "800 naira per tin" in replies[3]
        with app.state.svc.session_factory() as db:
            assert db.scalar(select(Customer).where(Customer.phone == CALLER)).language == "en"


def test_a_caller_who_just_carries_on_needs_no_answer(settings):
    """After the question, speaking clearly in a language is as good as naming it."""
    app, asked = app_with(settings, "en,yo")
    with TestClient(app) as client:
        replies = phone(client, ["[ha] Ina son madara", "[en] How much is Peak milk tin?"])
        assert "800 naira per tin" in replies[2]  # answered the price question straight away
        with app.state.svc.session_factory() as db:
            assert db.scalar(select(Customer).where(Customer.phone == CALLER)).language == "en"


def test_two_unclear_answers_and_she_carries_on_in_english_without_looping(settings):
    app, _ = app_with(settings, "en,yo")
    with TestClient(app) as client:
        replies = phone(client, ["[ha] blah blah", "[ha] blah blah again", "[ha] more blah", "[en] How much is Peak milk tin?"])
        assert replies[1] == "Sorry, I didn't catch your language. Would you like to speak English or Yoruba?"
        assert replies[2] == "Sorry, which language would you like to speak, English or Yoruba?"  # a different wording the 2nd time
        assert replies[3] == "Thank you. What would you like today?"  # gives up asking; English
        assert "800 naira per tin" in replies[4]
        with app.state.svc.session_factory() as db:
            assert db.scalar(select(Customer).where(Customer.phone == CALLER)).language is None  # a guess is not remembered


def test_the_question_is_also_asked_in_yoruba_when_we_have_the_wording(settings):
    app, _ = app_with(settings, "en,yo")
    templates.TRANSLATIONS["yo"] = {"language_ask": ["Yoruba wording: {languages}?"], "lang_name_en": ["Geesi"],
                                    "lang_name_yo": ["Yoruba"], "word_or": ["tabi"]}
    with TestClient(app) as client:
        start(client, "s2", CALLER)
        xml = raw_turn(client, "s2", CALLER, ("ha", "Ina son madara"))
        assert xml.count("<Play") == 2  # the English question, then the Yoruba one
        english = "Sorry, I didn't catch your language. Would you like to speak English or Yoruba?"
        yoruba = "Yoruba wording: Geesi tabi Yoruba?"
        assert AudioService.key(english, "en") in xml and AudioService.key(yoruba, "yo") in xml  # each in its own voice


def test_the_recording_notice_waits_until_the_language_is_settled(settings):
    app, _ = app_with(settings, "en,yo")
    with TestClient(app) as client:
        start(client, "s3", CALLER)
        asked_xml = raw_turn(client, "s3", CALLER, ("ha", "Ina son madara"))
        assert asked_xml.count("<Play") == 1  # just the question
        answered_xml = raw_turn(client, "s3", CALLER, ("yo", "Yoruba"))
        assert answered_xml.count("<Play") == 2  # recording notice, then "What would you like today?"


def test_there_is_no_key_press_route_any_more(settings):
    app, _ = app_with(settings, "en,yo")
    with TestClient(app) as client:
        assert client.post(f"/voice/language/{SECRET}", data={"sessionId": "x", "dtmfDigits": "2"}).status_code == 404


def test_a_language_that_was_switched_off_is_forgotten_for_returning_callers(settings):
    app, asked = app_with(settings, "en,yo")
    with app.state.svc.session_factory() as db:
        db.add(Customer(phone=CALLER, name="Musa", language="ha"))
        db.commit()
    with TestClient(app) as client:
        replies = phone(client, ["How much is Peak milk tin?"], lang="yo")
        assert "Welcome back, Musa" in replies[0]
        assert sorted(asked) == ["en", "yo"]  # detected afresh instead of trusting a disabled language
        with app.state.svc.session_factory() as db:
            assert db.scalar(select(Customer).where(Customer.phone == CALLER)).language == "yo"


# ---- the admin console -------------------------------------------------------------------------------------------


def test_merchant_language_must_be_an_enabled_one(client):
    login(client)
    form = {"name": "X", "owner_name": "Y", "owner_phone": "08031234568", "number": "+2348007776666"}
    bad = post(client, "/admin/merchants", {**form, "default_language": "ha"}, page="/admin/merchants")
    assert "language must be one of en, yo" in flash(client, bad)
    page = client.get("/admin/merchants").text
    assert '<option value="yo">' in page and '<option value="ha">' not in page and '<option value="ig">' not in page


def test_languages_page_and_dashboard_show_only_yoruba(client):
    login(client)
    page = client.get("/admin/languages").text
    assert "Yoruba" in page and "Hausa" not in page and "Igbo" not in page
    dash = client.get("/admin").text
    assert "Yoruba:" in dash and "Hausa:" not in dash


def test_label_page_offers_only_enabled_languages(client, call):
    call(["How much is Peak milk tin?"])
    login(client)
    html = client.get("/admin/label").text
    assert ">yo</option>" in html and ">ha</option>" not in html and ">ig</option>" not in html


def test_translation_sheet_has_only_the_yoruba_column(client, tmp_path):
    login(client)
    (tmp_path / "s.xlsx").write_bytes(client.get("/admin/translation-sheet.xlsx").content)
    header = [c.value for c in load_workbook(tmp_path / "s.xlsx")["Translate"][1]]
    assert header == ["key", "variant", "status", "when it is spoken", "English", "what {placeholders} will contain", "Yoruba", "notes"]


def test_full_sheet_export_is_still_available_for_later(tmp_path):
    export_xlsx(tmp_path / "all.xlsx")  # default: every language, for when Hausa and Igbo are switched on
    header = [c.value for c in load_workbook(tmp_path / "all.xlsx")["Translate"][1]]
    assert header[-4:] == ["Yoruba", "Hausa", "Igbo", "notes"]
    export_xlsx(tmp_path / "yo.xlsx", languages=("yo",))
    assert "Hausa" not in [c.value for c in load_workbook(tmp_path / "yo.xlsx")["Translate"][1]]
